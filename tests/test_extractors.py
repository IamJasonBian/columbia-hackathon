"""Local tests for the format extractors (no network; OCR tests skip without tesseract)."""

from __future__ import annotations

import io
from email.message import EmailMessage

import pytest
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont

from services.extraction.extractors import extract_document
from services.extraction.models import DocumentKind
from services.extraction.ocr import ocr_available

needs_ocr = pytest.mark.skipif(not ocr_available(), reason="tesseract not installed")


# ---------------------------------------------------------------------------
# Fixture builders
# ---------------------------------------------------------------------------

def make_text_pdf(lines: list[str]) -> bytes:
    """Hand-roll a one-page PDF with a real text layer (Helvetica)."""
    stream = "BT /F1 12 Tf 72 720 Td 14 TL " + " ".join(
        f"({line}) Tj T*" for line in lines
    ) + " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        "/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n{body}\nendobj\n".encode())
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Deductions"
    ws.append(["Invoice", "Date", "Amount"])
    ws.append(["INV-1", "2026-09-01", -125.5])
    ws.append(["INV-2", "2026-09-02", 300])
    summary = wb.create_sheet("Summary")
    summary.append(["Total", 174.5])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_scan_image(text: str, fmt: str = "TIFF") -> bytes:
    image = Image.new("L", (1200, 300), 255)
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 48)
    except OSError:
        font = ImageFont.load_default(size=48)
    draw.text((40, 100), text, fill=0, font=font)
    buf = io.BytesIO()
    image.save(buf, format=fmt, dpi=(300, 300))
    return buf.getvalue()


def make_eml(attachments: list[tuple[str, bytes, str, str]]) -> bytes:
    msg = EmailMessage()
    msg["From"] = "Billing <billing@example.com>"
    msg["To"] = "ops@example.com, ap@example.com"
    msg["Subject"] = "Remittance for Check #: 12345"
    msg["Date"] = "Fri, 09 Oct 2026 10:00:00 -0400"
    msg["Message-ID"] = "<abc123@example.com>"
    msg.set_content("Please find the remittance attached.\nThanks")
    msg.add_alternative("<p>Please find the <b>remittance</b> attached.</p>", subtype="html")
    for filename, data, maintype, subtype in attachments:
        msg.add_attachment(data, maintype=maintype, subtype=subtype, filename=filename)
    return msg.as_bytes()


# ---------------------------------------------------------------------------
# Detection
# ---------------------------------------------------------------------------

def test_pdf_text_layer():
    doc = extract_document(make_text_pdf(["Invoice 14210", "Total 3217.50"]), "r.pdf")
    assert doc.kind == DocumentKind.pdf
    assert doc.metadata["page_count"] == 1
    assert "Invoice 14210" in doc.text and "3217.50" in doc.text
    assert not doc.pages[0].ocr
    assert doc.sha256 and doc.size_bytes > 0


@needs_ocr
def test_scanned_pdf_falls_back_to_ocr():
    image = Image.open(io.BytesIO(make_scan_image("SHORTAGE 4210")))
    buf = io.BytesIO()
    image.save(buf, format="PDF", resolution=300)
    doc = extract_document(buf.getvalue(), "scan.pdf")
    assert doc.pages[0].ocr
    assert "4210" in doc.text


# ---------------------------------------------------------------------------
# Spreadsheets
# ---------------------------------------------------------------------------

def test_xlsx_every_sheet_with_records():
    doc = extract_document(make_xlsx(), "d.xlsx")
    assert [t.name for t in doc.tables] == ["Deductions", "Summary"]
    deductions = doc.tables[0]
    assert deductions.headers == ["Invoice", "Date", "Amount"]
    assert deductions.records[0] == {"Invoice": "INV-1", "Date": "2026-09-01", "Amount": -125.5}
    assert "INV-2" in doc.text


def test_email_headers_body_and_recursive_attachments():
    eml = make_eml(
        [
            ("remit.pdf", make_text_pdf(["Check 12345"]), "application", "pdf"),
            ("lines.csv", b"invoice,net\n14210C1,-3217.50\n", "text", "csv"),
            ("detail.xlsx", make_xlsx(), "application",
             "vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        ]
    )
    doc = extract_document(eml, "msg.eml")
    assert doc.kind == DocumentKind.email
    assert doc.metadata["subject"] == "Remittance for Check #: 12345"
    assert doc.metadata["to"] == ["ops@example.com", "ap@example.com"]
    assert doc.metadata["message_id"] == "abc123@example.com"
    assert doc.metadata["date"].startswith("2026-10-09")
    assert doc.text.startswith("Please find the remittance attached.")
    kinds = {a.filename: a.kind for a in doc.attachments}
    assert kinds == {
        "remit.pdf": DocumentKind.pdf,
        "lines.csv": DocumentKind.spreadsheet,
        "detail.xlsx": DocumentKind.spreadsheet,
    }
    csv_doc = next(a for a in doc.attachments if a.filename == "lines.csv")
    assert csv_doc.tables[0].records == [{"invoice": "14210C1", "net": "-3217.50"}]
