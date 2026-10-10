"""Local tests for extraction_service (no network; OCR tests skip without tesseract)."""

from __future__ import annotations

import io
import json
import os
from email.message import EmailMessage
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont

os.environ.setdefault("SKIP_AUTH", "1")

from extraction_service.extractors import detect_kind, extract_document  # noqa: E402
from extraction_service.llm import extract_structured, render_document_for_prompt  # noqa: E402
from extraction_service.main import app  # noqa: E402
from extraction_service.models import DocumentKind  # noqa: E402
from extraction_service.ocr import ocr_available  # noqa: E402

needs_ocr = pytest.mark.skipif(not ocr_available(), reason="tesseract not installed")


# ---------------------------------------------------------------------------
# Fixture builders
# ---------------------------------------------------------------------------

def make_text_pdf(lines: list[str]) -> bytes:
    """Hand-roll a one-page PDF with a real text layer (Helvetica)."""
    stream = "BT /F1 12 Tf 72 720 Td 14 TL " + " ".join(
        f"({line}) Tj T*" for line in lines
    ) + " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        "/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n{body}\nendobj\n".encode())
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Deductions"
    ws.append(["Invoice", "Date", "Amount"])
    ws.append(["INV-1", "2026-09-01", -125.5])
    ws.append(["INV-2", "2026-09-02", 300])
    summary = wb.create_sheet("Summary")
    summary.append(["Total", 174.5])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_scan_image(text: str, fmt: str = "TIFF") -> bytes:
    image = Image.new("L", (1200, 300), 255)
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 48)
    except OSError:
        font = ImageFont.load_default(size=48)
    draw.text((40, 100), text, fill=0, font=font)
    buf = io.BytesIO()
    image.save(buf, format=fmt, dpi=(300, 300))
    return buf.getvalue()


def make_eml(attachments: list[tuple[str, bytes, str, str]]) -> bytes:
    msg = EmailMessage()
    msg["From"] = "Billing <billing@example.com>"
    msg["To"] = "ops@example.com, ap@example.com"
    msg["Subject"] = "Remittance for Check #: 12345"
    msg["Date"] = "Fri, 09 Oct 2026 10:00:00 -0400"
    msg["Message-ID"] = "<abc123@example.com>"
    msg.set_content("Please find the remittance attached.\nThanks")
    msg.add_alternative("<p>Please find the <b>remittance</b> attached.</p>", subtype="html")
    for filename, data, maintype, subtype in attachments:
        msg.add_attachment(data, maintype=maintype, subtype=subtype, filename=filename)
    return msg.as_bytes()


# ---------------------------------------------------------------------------
# Detection
# ---------------------------------------------------------------------------

def test_detect_kind():
    assert detect_kind("x.bin", make_text_pdf(["hi"])) == DocumentKind.pdf
    assert detect_kind("a.xlsx", make_xlsx()) == DocumentKind.spreadsheet
    assert detect_kind("a.csv", b"a,b\n1,2\n") == DocumentKind.spreadsheet
    assert detect_kind("scan", make_scan_image("x")) == DocumentKind.image
    assert detect_kind("noext", make_eml([])) == DocumentKind.email
    assert detect_kind("notes.txt", b"hello") == DocumentKind.text
    assert detect_kind("blob", b"\x00\x01") == DocumentKind.unknown


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def test_pdf_text_layer():
    doc = extract_document(make_text_pdf(["Invoice 14210", "Total 3217.50"]), "r.pdf")
    assert doc.kind == DocumentKind.pdf
    assert doc.metadata["page_count"] == 1
    assert "Invoice 14210" in doc.text and "3217.50" in doc.text
    assert not doc.pages[0].ocr
    assert doc.sha256 and doc.size_bytes > 0


@needs_ocr
def test_scanned_pdf_falls_back_to_ocr():
    image = Image.open(io.BytesIO(make_scan_image("SHORTAGE 4210")))
    buf = io.BytesIO()
    image.save(buf, format="PDF", resolution=300)
    doc = extract_document(buf.getvalue(), "scan.pdf")
    assert doc.pages[0].ocr
    assert "4210" in doc.text


# ---------------------------------------------------------------------------
# Spreadsheets
# ---------------------------------------------------------------------------

def test_xlsx_every_sheet_with_records():
    doc = extract_document(make_xlsx(), "d.xlsx")
    assert [t.name for t in doc.tables] == ["Deductions", "Summary"]
    deductions = doc.tables[0]
    assert deductions.headers == ["Invoice", "Date", "Amount"]
    assert deductions.records[0] == {"Invoice": "INV-1", "Date": "2026-09-01", "Amount": -125.5}
    assert "INV-2" in doc.text


def test_csv_sniffs_delimiter_and_strips_bom():
    data = "﻿sku;qty\nA1;3\n\nB2;5\n".encode("utf-8")
    doc = extract_document(data, "items.csv")
    table = doc.tables[0]
    assert table.headers == ["sku", "qty"]
    assert table.rows == [["A1", "3"], ["B2", "5"]]


def test_legacy_xls_warns():
    doc = extract_document(b"\xd0\xcf\x11\xe0", "old.xls")
    assert doc.warnings and "xls" in doc.warnings[0]


# ---------------------------------------------------------------------------
# Email
# ---------------------------------------------------------------------------

def test_email_headers_body_and_recursive_attachments():
    eml = make_eml(
        [
            ("remit.pdf", make_text_pdf(["Check 12345"]), "application", "pdf"),
            ("lines.csv", b"invoice,net\n14210C1,-3217.50\n", "text", "csv"),
            ("detail.xlsx", make_xlsx(), "application",
             "vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        ]
    )
    doc = extract_document(eml, "msg.eml")
    assert doc.kind == DocumentKind.email
    assert doc.metadata["subject"] == "Remittance for Check #: 12345"
    assert doc.metadata["to"] == ["ops@example.com", "ap@example.com"]
    assert doc.metadata["message_id"] == "abc123@example.com"
    assert doc.metadata["date"].startswith("2026-10-09")
    assert doc.text.startswith("Please find the remittance attached.")
    kinds = {a.filename: a.kind for a in doc.attachments}
    assert kinds == {
        "remit.pdf": DocumentKind.pdf,
        "lines.csv": DocumentKind.spreadsheet,
        "detail.xlsx": DocumentKind.spreadsheet,
    }
    csv_doc = next(a for a in doc.attachments if a.filename == "lines.csv")
    assert csv_doc.tables[0].records == [{"invoice": "14210C1", "net": "-3217.50"}]


def test_html_only_email_is_flattened():
    msg = EmailMessage()
    msg["Subject"] = "html"
    msg["From"] = "a@example.com"
    msg.set_content("<html><style>x{}</style><p>Line one</p><p>Line <b>two</b></p></html>", subtype="html")
    doc = extract_document(msg.as_bytes(), "h.eml")
    assert doc.text == "Line one\nLine two"


@needs_ocr
def test_email_with_scanned_tiff_attachment():
    eml = make_eml([("scan.tif", make_scan_image("CHECK 8500"), "image", "tiff")])
    doc = extract_document(eml, "m.eml")
    assert "8500" in doc.attachments[0].text


# ---------------------------------------------------------------------------
# OpenAI layer (fake client — no network)
# ---------------------------------------------------------------------------

class FakeResponses:
    def __init__(self, payload: dict):
        self.payload = payload
        self.calls: list[dict] = []

    def create(self, **kwargs):
        self.calls.append(kwargs)
        return SimpleNamespace(output_text=json.dumps(self.payload), usage=None)


def test_render_document_for_prompt_includes_attachments():
    eml = make_eml([("lines.csv", b"invoice,net\n14210C1,-3217.50\n", "text", "csv")])
    text = render_document_for_prompt(extract_document(eml, "m.eml"))
    assert "=== EMAIL: m.eml ===" in text
    assert "--- sheet: lines.csv ---" in text
    assert "14210C1\t-3217.50" in text


def test_extract_structured_with_schema():
    schema = {
        "type": "object",
        "properties": {"check_number": {"type": "string"}},
        "required": ["check_number"],
        "additionalProperties": False,
    }
    fake = SimpleNamespace(responses=FakeResponses({"check_number": "12345"}))
    doc = extract_document(make_text_pdf(["Check 12345"]), "r.pdf")
    result = extract_structured(doc, schema=schema, strict=True, client=fake, model="m")
    assert result.data == {"check_number": "12345"}
    call = fake.responses.calls[0]
    assert call["text"]["format"]["type"] == "json_schema"
    assert call["text"]["format"]["strict"] is True
    assert "Check 12345" in call["input"][1]["content"][0]["text"]


def test_extract_structured_free_form_json():
    fake = SimpleNamespace(responses=FakeResponses({"total": 174.5}))
    doc = extract_document(make_xlsx(), "d.xlsx")
    result = extract_structured(doc, instructions="Sum the amounts", client=fake, model="m")
    assert result.data == {"total": 174.5}
    call = fake.responses.calls[0]
    assert call["text"]["format"] == {"type": "json_object"}
    assert "Sum the amounts" in call["input"][0]["content"]


# ---------------------------------------------------------------------------
# HTTP API
# ---------------------------------------------------------------------------

def test_api_extract_and_crud():
    client = TestClient(app)
    assert client.get("/health").json()["service"] == "extraction"

    resp = client.post("/extract", files={"file": ("d.xlsx", make_xlsx())})
    assert resp.status_code == 200
    row = resp.json()
    assert row["kind"] == "spreadsheet"
    assert row["document"]["tables"][0]["records"][1]["Invoice"] == "INV-2"

    pdf = client.post("/extract/pdf", files={"file": ("r.pdf", make_text_pdf(["Hello PDF"]))})
    assert "Hello PDF" in pdf.json()["document"]["text"]

    eml = client.post("/extract/email", files={"file": ("m.eml", make_eml([]))})
    assert eml.json()["document"]["metadata"]["subject"].startswith("Remittance")

    assert client.get(f"/extractions/{row['id']}").json()["id"] == row["id"]
    assert {r["kind"] for r in client.get("/extractions?kind=spreadsheet").json()} == {"spreadsheet"}
    assert client.delete(f"/extractions/{row['id']}").json() == {"deleted": True}
    assert client.get(f"/extractions/{row['id']}").status_code == 404


def test_api_rejects_empty_upload_and_bad_schema():
    client = TestClient(app)
    assert client.post("/extract", files={"file": ("e.pdf", b"")}).status_code == 400
    bad = client.post(
        "/extract/structured",
        files={"file": ("a.csv", b"a,b\n1,2\n")},
        data={"json_schema": "{not json"},
    )
    assert bad.status_code == 400
