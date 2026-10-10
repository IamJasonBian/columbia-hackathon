"""Spreadsheet extraction: .xlsx/.xlsm (openpyxl) and .csv/.tsv (stdlib).

Each sheet becomes one ExtractedTable whose header row is the first
non-empty row. Cached formula values are read (``data_only=True``), so a
workbook saved by Excel yields numbers, not ``=SUM(...)`` strings.
"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime, time
from typing import Any

from extraction_service.models import DocumentKind, ExtractedDocument, ExtractedTable

XLSX_EXTENSIONS = (".xlsx", ".xlsm", ".xltx", ".xltm")
CSV_EXTENSIONS = (".csv", ".tsv", ".txt")


def _normalize(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, str):
        return value.strip()
    return value


def _table_from_rows(name: str, raw_rows: list[list[Any]]) -> ExtractedTable:
    rows = [[_normalize(v) for v in row] for row in raw_rows]
    rows = [row for row in rows if any(v not in (None, "") for v in row)]
    if not rows:
        return ExtractedTable(name=name)
    width = max(len(row) for row in rows)
    rows = [row + [None] * (width - len(row)) for row in rows]
    # Trim columns that are empty in every row (openpyxl reports the used
    # range, which often includes formatted-but-blank trailing columns).
    keep = [i for i in range(width) if any(row[i] not in (None, "") for row in rows)]
    rows = [[row[i] for i in keep] for row in rows]
    headers = [
        str(h) if h not in (None, "") else f"col_{i}" for i, h in enumerate(rows[0])
    ]
    return ExtractedTable(name=name, headers=headers, rows=rows[1:])


def _table_text(table: ExtractedTable) -> str:
    lines = [f"# {table.name}", "\t".join(table.headers)]
    lines += ["\t".join("" if v is None else str(v) for v in row) for row in table.rows]
    return "\n".join(lines)


def extract_xlsx(data: bytes, filename: str) -> ExtractedDocument:
    """Extract every worksheet of an Excel workbook.

    Args:
        data: Raw .xlsx bytes.
        filename: Original filename.

    Returns:
        The extracted document, one table per sheet.
    """
    from openpyxl import load_workbook

    doc = ExtractedDocument(
        filename=filename,
        kind=DocumentKind.spreadsheet,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        doc.warnings.append(f"could not open workbook: {e}")
        return doc
    try:
        for sheet in workbook.worksheets:
            raw_rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            doc.tables.append(_table_from_rows(sheet.title, raw_rows))
    finally:
        workbook.close()
    doc.metadata = {"sheets": [t.name for t in doc.tables]}
    doc.text = "\n\n".join(_table_text(t) for t in doc.tables)
    return doc


def extract_csv(data: bytes, filename: str) -> ExtractedDocument:
    """Extract a delimited text file, sniffing the delimiter.

    Args:
        data: Raw file bytes (UTF-8, UTF-8-BOM, or Latin-1).
        filename: Original filename.

    Returns:
        The extracted document with a single table.
    """
    doc = ExtractedDocument(filename=filename, kind=DocumentKind.spreadsheet, content_type="text/csv")
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
        doc.warnings.append("not valid UTF-8; decoded as Latin-1")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel_tab if filename.lower().endswith(".tsv") else csv.excel
    raw_rows = [list(row) for row in csv.reader(io.StringIO(text), dialect)]
    table = _table_from_rows(filename, raw_rows)
    doc.tables.append(table)
    doc.text = _table_text(table)
    return doc


def extract_spreadsheet(data: bytes, filename: str) -> ExtractedDocument:
    """Route a spreadsheet to the xlsx or csv extractor.

    Args:
        data: Raw file bytes.
        filename: Original filename (extension picks the reader).

    Returns:
        The extracted document.
    """
    lower = filename.lower()
    if lower.endswith(XLSX_EXTENSIONS) or data[:4] == b"PK\x03\x04":
        return extract_xlsx(data, filename)
    if lower.endswith(".xls"):
        doc = ExtractedDocument(filename=filename, kind=DocumentKind.spreadsheet)
        doc.warnings.append("legacy .xls is not supported; re-save as .xlsx or .csv")
        return doc
    return extract_csv(data, filename)
